"""Rule-based checks for claims and fact-registry term mentions in application documents.

This module does not establish that an entire document is true. It reports only
claims found by the configured patterns and whether the fact registry returned
evidence IDs for those claims.
"""

import json
import re
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal, InvalidOperation
from pathlib import Path

from src.fact_registry import get_fact_registry


ROOT = Path(__file__).resolve().parents[2]

_METRIC_PATTERN = re.compile(
    r"(?:-?\d+(?:[.,]\d+)?\s*(?:kcal/mol|µM|μM|nM|mM|mg/ml|mg/ml|fold|%|p\s*[=<])"
    r"|PDB\s*[A-Z0-9]+|-11\.9|-9\.4|0\.007|3ODU)",
    re.IGNORECASE,
)
_DISCOVERY_PATTERN = re.compile(
    r"\b(lead candidate|novel derivative|statistically significant|improved survival|"
    r"differentially expressed|binding affinity)\b",
    re.IGNORECASE,
)
_SKIP_PREFIX = re.compile(r"^(phone|\+90|istanbul|2021|2025|2027)\b", re.IGNORECASE)
_NUMBER_TOKEN = re.compile(r"(?<![A-Za-z0-9])[-+]?\d+(?:[.,]\d+)?(?![A-Za-z0-9])")
_PDB_ID = re.compile(r"\bPDB\s*[A-Z0-9]+\b", re.IGNORECASE)
_STOP_WORDS = {"a", "an", "and", "at", "by", "for", "from", "in", "of", "on", "the", "to", "vs", "with"}


def _scientific_values(text):
    """Return normalized numeric values, excluding years and PDB structure IDs."""
    cleaned = _PDB_ID.sub(" ", text or "")
    values = set()
    for token in _NUMBER_TOKEN.findall(cleaned):
        normalized = token.replace(",", ".")
        try:
            value = Decimal(normalized)
        except InvalidOperation:
            continue
        if value == value.to_integral_value() and 1900 <= value <= 2100:
            continue
        values.add(value.normalize())
    return values


def _has_supporting_phrase(claim, fact_text):
    """Require an exact scientific phrase, not just two coincidentally shared words."""
    claim_words = [word.lower() for word in re.findall(r"[a-zA-Z0-9]+", claim)]
    fact_words = [word.lower() for word in re.findall(r"[a-zA-Z0-9]+", fact_text)]
    if not claim_words or not fact_words:
        return False

    # Named discovery patterns are meaningful two-word phrases and must occur verbatim.
    discovery_phrases = [match.group(0).lower() for match in _DISCOVERY_PATTERN.finditer(claim)]
    fact_lower = fact_text.lower()
    if any(phrase in fact_lower for phrase in discovery_phrases):
        return True

    # Otherwise require a contiguous three-token phrase after removing common glue words.
    claim_content = [word for word in claim_words if word not in _STOP_WORDS]
    fact_content = [word for word in fact_words if word not in _STOP_WORDS]
    if len(claim_content) < 3:
        return False
    fact_trigrams = {
        tuple(fact_content[index:index + 3])
        for index in range(len(fact_content) - 2)
    }
    return any(
        tuple(claim_content[index:index + 3]) in fact_trigrams
        for index in range(len(claim_content) - 2)
    )


def _filter_supported_fact_ids(claim, candidate_ids, fact_registry):
    """Keep only IDs that collectively support all claim values and a concrete phrase."""
    claim_values = _scientific_values(claim)
    supported = []
    supported_values = set()
    for fact_id in candidate_ids:
        fact = fact_registry.get(fact_id)
        if fact is None or not _has_supporting_phrase(claim, fact.text):
            continue
        values = _scientific_values(fact.text)
        values.update(Decimal(str(value)).normalize() for value in fact.metrics)
        if claim_values and not claim_values.intersection(values):
            continue
        supported.append(fact_id)
        supported_values.update(values)

    if claim_values and not claim_values.issubset(supported_values):
        return []
    return supported


def extract_text_from_docx(docx_path):
    """Extract text from a DOCX file; raise when the artifact cannot be read."""
    try:
        from docx import Document

        doc = Document(docx_path)
        texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        texts.append(cell.text.strip())
        return " ".join(texts)
    except Exception:
        try:
            with zipfile.ZipFile(docx_path) as archive:
                xml_content = archive.read("word/document.xml")
            tree = ET.fromstring(xml_content)
            return " ".join(node.text for node in tree.iter() if node.text)
        except Exception as fallback_error:
            raise ValueError(f"Unable to read DOCX artifact: {docx_path}") from fallback_error


def normalize_number(num_str):
    """Normalize string numbers to float for epsilon comparison."""
    try:
        return float(num_str.replace(",", "."))
    except (ValueError, AttributeError):
        return None


def extract_numbers_from_text(text):
    """Extract numeric values, omitting common document date years."""
    nums = []
    for raw in re.findall(r"-?\d+(?:[.,]\d+)?", text):
        value = normalize_number(raw)
        if value is not None and abs(value) not in [2021, 2025, 2026, 2027]:
            nums.append(value)
    return nums


def extract_hard_scientific_claims(text, master_cv=None):
    """Return every pattern-detected claim, with registry evidence IDs when found.

    Detection is deliberately rule-based and incomplete. Every detected sentence
    is assessed; absence from this list does not imply that a sentence is true.
    """
    cleaned = re.sub(r"\s+", " ", text or "")
    sentences = re.split(r"(?<=[.!?])\s+", cleaned)
    critical_claims = []
    fact_registry = get_fact_registry()

    for sentence in sentences:
        claim = sentence.strip()
        if len(claim) <= 25 or _SKIP_PREFIX.match(claim):
            continue

        has_metric = bool(_METRIC_PATTERN.search(claim))
        has_discovery = bool(_DISCOVERY_PATTERN.search(claim))
        if not (has_metric or has_discovery):
            continue
        if any(existing["claim"] == claim for existing in critical_claims):
            continue

        candidate_fact_ids = fact_registry.match_claim_to_facts(claim)
        matched_fact_ids = _filter_supported_fact_ids(claim, candidate_fact_ids, fact_registry)
        claim_lower = claim.lower()
        target_mismatch = "-11.9" in claim and not (
            "cxcr4" in claim_lower or "amentoflavone" in claim_lower
        )
        if target_mismatch:
            status = "FLAGGED_CONTEXT_MISMATCH"
            reference = "The -11.9 metric was associated with CXCR4/Amentoflavone in the fact registry."
        elif matched_fact_ids:
            status = "VERIFIED_ACCURATE"
            matched_facts = [fact_registry.get(fid) for fid in matched_fact_ids]
            reference = " | ".join(fact.text for fact in matched_facts if fact is not None)
        else:
            status = "FLAGGED_UNVERIFIED_CLAIM"
            reference = "No matching fact-registry evidence ID was found."

        critical_claims.append({
            "claim": claim,
            "type": "QUANTITATIVE_METRIC" if has_metric else "SCIENTIFIC_FINDING",
            "status": status,
            "evidence_fact_ids": matched_fact_ids,
            "master_reference": reference,
        })

    return critical_claims


def _read_artifact_text(folder, filename, *, json_file=False):
    path = folder / filename
    try:
        if not path.is_file():
            return "", {"readable": False, "reason": "missing"}
        if json_file:
            content = path.read_text(encoding="utf-8")
            json.loads(content)
            return content, {"readable": True, "reason": None}
        return extract_text_from_docx(path), {"readable": True, "reason": None}
    except (OSError, UnicodeError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        return "", {"readable": False, "reason": type(exc).__name__}


def verify_tailored_document(folder_path, master_cv):
    """Create a v3 report with explicit scope and complete coverage of detected claims."""
    folder = Path(folder_path)
    fact_registry = get_fact_registry()
    artifacts = {}
    content = []
    for filename, is_json in (
        ("tailored_cv.docx", False),
        ("cover_letter.docx", False),
        ("apply_info.json", True),
    ):
        text, state = _read_artifact_text(folder, filename, json_file=is_json)
        artifacts[filename] = state
        content.append(text)

    doc_text = " ".join(content)
    doc_text_lower = doc_text.lower()

    # These are detected mentions of registry terms, not truth judgments.
    detected_terms = []
    for fact in fact_registry.facts.values():
        for entity in fact.entities:
            entity = entity.strip()
            if entity and entity.lower() not in {term.lower() for term in detected_terms}:
                if re.search(r"\b" + re.escape(entity) + r"\b", doc_text_lower, re.IGNORECASE):
                    detected_terms.append(entity)

    claims = extract_hard_scientific_claims(doc_text, master_cv)
    unsupported_count = sum(claim["status"] != "VERIFIED_ACCURATE" for claim in claims)
    all_artifacts_readable = all(state["readable"] for state in artifacts.values())
    claim_scope = (
        "No scientific claims matched the configured detection patterns; this does not establish full-document truth."
        if not claims else
        "Only sentences matching the configured quantitative/scientific patterns are assessed; unrecognized claims may be missed."
    )
    checks = {
        # Retained as informational indicators for older consumers. They are not universal requirements.
        "institution_preserved": any(term in doc_text_lower for term in ("example research lab", "example university", "example university")),
        "target_gene_accurate": "cxcr4" in doc_text_lower or "protein" in doc_text_lower,
        "flagged_claims_count": len(claims),
        "unsupported_claims_count": unsupported_count,
        "artifacts_readable": all_artifacts_readable,
    }
    report = {
        "version": "3.0.0",
        "term_fidelity_score": None,
        "registered_terms_detected_count": len(detected_terms),
        "registered_terms_detected": detected_terms,
        "artifact_read_status": artifacts,
        "claims_with_references": claims,
        "detected_claim_coverage": {
            "detection_method": "Configured regular expressions for quantitative and selected scientific phrases",
            "detected_claims_count": len(claims),
            "assessed_claims_count": len(claims) if all_artifacts_readable else 0,
            "coverage_status": "complete" if all_artifacts_readable else "incomplete",
            "scope_and_limitations": claim_scope,
        },
        "factual_checks": checks,
        "audit_disclaimer": (
            "Rule-based partial audit only. A VERIFIED_ACCURATE status means the registry matcher returned the listed "
            "fact IDs; this report does not certify every document statement or overall truth."
        ),
    }

    (folder / "scientific_validation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return report
